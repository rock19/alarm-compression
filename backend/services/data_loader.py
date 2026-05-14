from __future__ import annotations
import openpyxl
from datetime import datetime
from typing import Optional

REQUIRED_COLUMNS = ["网元", "告警名称", "告警级别", "发生时间"]

# Column name aliases: standard name → accepted variations
COLUMN_ALIASES = {
    "发生时间": ["发生时间", "首次发生时间", "告警发生时间", "发生时间(首次)", "开始时间"],
    "恢复时间": ["恢复时间", "最后发生时间", "清除时间", "恢复时间(最后)"],
    "铁路线": ["铁路线", "线路"],
    "恢复状态": ["恢复状态", "告警状态"],
    "告警编码": ["告警编码", "告警ID"],
}

# Legacy position-based map (fallback)
COLUMN_MAP = {
    0: "专业", 1: "网管", 2: "网元", 3: "告警对象", 4: "告警级别",
    5: "告警名称", 6: "告警类型", 7: "告警描述", 8: "发生时间",
    9: "恢复时间", 10: "恢复状态", 11: "关联业务", 12: "组织机构",
    13: "确认状态", 14: "确认时间", 15: "确认人", 16: "告警有效性",
    17: "无效原因", 18: "业务使用单位", 19: "标记人", 20: "标记时间",
    21: "障碍诊断", 22: "告警分析", 23: "告警编码", 24: "子系统确认状态",
    25: "子系统确认时间", 26: "清除时间", 27: "清除人", 28: "铁路线",
    29: "站点", 30: "机房", 31: "厂商", 32: "告警标识",
}


def _build_header_index(headers: list) -> dict[str, int]:
    """Build header_name → column_index from actual file headers."""
    index: dict[str, int] = {}
    for i, h in enumerate(headers):
        if h is None:
            continue
        h_clean = str(h).strip()
        index[h_clean] = i
    return index


def _resolve_column(header_index: dict[str, int], row: tuple, standard_name: str, col_idx: int):
    """
    Get value for a standard column name using header-aware lookup.
    1. Try exact match against standard_name
    2. Try alias matches
    3. Fall back to position-based COLUMN_MAP
    """
    # Try exact match first
    if standard_name in header_index:
        i = header_index[standard_name]
        if i < len(row):
            return row[i]

    # Try aliases
    aliases = COLUMN_ALIASES.get(standard_name, [standard_name])
    for alias in aliases:
        if alias in header_index:
            i = header_index[alias]
            if i < len(row):
                val = row[i]
                # For time columns, ensure we return something usable
                if standard_name == "发生时间" and isinstance(val, (int, float)):
                    continue  # Skip numeric values (like 告警次数), try next alias
                return val

    # Fallback to position-based
    if col_idx < len(row):
        return row[col_idx]

    return None


def load_excel(file_path: str) -> tuple[list[dict], dict]:
    """Load alarm Excel file. Auto-detects column format from headers."""
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    header_index = _build_header_index(headers)

    # Standard column positions for fallback
    STD_COLS = [
        ("专业", 0), ("网管", 1), ("网元", 2), ("告警对象", 3), ("告警级别", 4),
        ("告警名称", 5), ("告警类型", 6), ("告警描述", 7), ("发生时间", 8),
        ("恢复时间", 9), ("恢复状态", 10), ("关联业务", 11), ("组织机构", 12),
        ("确认状态", 13), ("确认时间", 14), ("确认人", 15), ("告警有效性", 16),
        ("无效原因", 17), ("业务使用单位", 18), ("标记人", 19), ("标记时间", 20),
        ("障碍诊断", 21), ("告警分析", 22), ("告警编码", 23), ("子系统确认状态", 24),
        ("子系统确认时间", 25), ("清除时间", 26), ("清除人", 27), ("铁路线", 28),
        ("站点", 29), ("机房", 30), ("厂商", 31), ("告警标识", 32),
    ]

    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        record = {}
        for std_name, fallback_idx in STD_COLS:
            record[std_name] = _resolve_column(header_index, row, std_name, fallback_idx)

        errors = validate_alarm_record(record)
        if not errors:
            record["发生时间"] = _parse_datetime(record.get("发生时间"))
            record["恢复时间"] = _parse_datetime(record.get("恢复时间"))
            records.append(record)

    wb.close()

    meta = {
        "total_records": len(records),
        "unique_ne": len(set(r["网元"] for r in records if r["网元"])),
        "unique_alarm_names": len(set(r["告警名称"] for r in records if r["告警名称"])),
        "time_min": min((r["发生时间"] for r in records if r["发生时间"]), default=None),
        "time_max": max((r["发生时间"] for r in records if r["发生时间"]), default=None),
        "railway_lines": list(set(r["铁路线"] for r in records if r["铁路线"])),
        "severity_levels": list(set(r["告警级别"] for r in records if r["告警级别"])),
    }
    return records, meta


def validate_alarm_record(record: dict) -> list[str]:
    """Validate a single alarm record. Returns list of error messages."""
    errors = []
    for col in REQUIRED_COLUMNS:
        if col not in record or record[col] is None or str(record[col]).strip() == "":
            errors.append(f"缺少必填列: {col}")
    if "发生时间" in record and record["发生时间"]:
        dt = _parse_datetime(record["发生时间"])
        if dt is None:
            errors.append(f"发生时间格式无效: {record['发生时间']}")
    return errors


def _parse_datetime(val) -> Optional[datetime]:
    if val is None:
        return None
    if isinstance(val, datetime):
        return val
    for fmt in [
        "%Y-%m-%d %H:%M:%S.%f", "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M", "%Y/%m/%d %H:%M:%S",
        "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M:%S.%f",
    ]:
        try:
            return datetime.strptime(str(val).strip(), fmt)
        except ValueError:
            continue
    return None


def clean_continuous_alarms(records: list[dict]) -> list[dict]:
    """
    Remove continuous duplicate alarms from the same source.

    Two alarms with identical (专业, 网管, 网元, 告警对象, 告警级别,
    告警名称, 告警类型, 告警描述) are considered duplicates if the
    second alarm's 发生时间 falls within [first.发生时间, first.恢复时间].

    Returns records sorted by 发生时间.
    """

    DEDUP_FIELDS = [
        "专业", "网管", "网元", "告警对象", "告警级别",
        "告警名称", "告警类型", "告警描述",
    ]

    groups: dict[tuple, list[dict]] = {}
    for r in records:
        key = tuple(r.get(f, "") or "" for f in DEDUP_FIELDS)
        groups.setdefault(key, []).append(r)

    kept = []
    for key, group in groups.items():
        group.sort(key=lambda r: r.get("发生时间") or datetime.min)
        base = None
        for r in group:
            occ = r.get("发生时间")
            rec = r.get("恢复时间")
            if occ is None:
                kept.append(r)
                continue
            if base is None:
                base = r
                kept.append(r)
                continue
            base_rec = base.get("恢复时间")
            if base_rec is not None and occ <= base_rec:
                continue
            base = r
            kept.append(r)

    kept.sort(key=lambda r: r.get("发生时间") or datetime.min)
    return kept


def load_excel_multiple(file_paths: list[str]) -> tuple[list[dict], dict]:
    """
    Load and merge multiple alarm Excel files, clean duplicate continuous alarms.
    Files are loaded individually, merged, sorted by 发生时间, then cleaned.
    """
    from datetime import datetime

    all_records = []
    for path in file_paths:
        records, _ = load_excel(path)
        all_records.extend(records)

    all_records.sort(key=lambda r: r.get("发生时间") or datetime.min)
    cleaned = clean_continuous_alarms(all_records)

    meta = {
        "total_records": len(cleaned),
        "unique_ne": len(set(r["网元"] for r in cleaned if r.get("网元"))),
        "unique_alarm_names": len(set(r["告警名称"] for r in cleaned if r.get("告警名称"))),
        "time_min": min((r["发生时间"] for r in cleaned if r.get("发生时间")), default=None),
        "time_max": max((r["发生时间"] for r in cleaned if r.get("发生时间")), default=None),
        "railway_lines": list(set(r.get("铁路线", "") for r in cleaned if r.get("铁路线"))),
        "severity_levels": list(set(r.get("告警级别", "") for r in cleaned if r.get("告警级别"))),
    }
    return cleaned, meta
