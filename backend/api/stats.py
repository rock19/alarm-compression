from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from services.store import store
from schemas.schemas import StatsResponse
from collections import Counter
import io
import openpyxl
from datetime import datetime

router = APIRouter()


@router.get("/stats", response_model=StatsResponse)
async def get_stats():
    alarms = store.get_alarms()
    meta = store.get_meta()

    if not alarms:
        raise HTTPException(status_code=400, detail="请先上传告警数据")

    severity_dist = dict(Counter(a["告警级别"] for a in alarms if a.get("告警级别")))
    alarm_name_dist = Counter(a["告警名称"] for a in alarms if a.get("告警名称"))
    ne_dist = Counter(a["网元"] for a in alarms if a.get("网元"))
    railway_dist = dict(Counter(a["铁路线"] for a in alarms if a.get("铁路线")))

    time_min = meta.get("time_min")
    time_max = meta.get("time_max")

    # Build hourly time-series for all alarms
    hour_buckets: dict[str, int] = {}
    for a in alarms:
        t = a.get("发生时间")
        if t:
            key = t.strftime("%Y-%m-%dT%H:00:00")
            hour_buckets[key] = hour_buckets.get(key, 0) + 1

    time_series = [{"time": k, "count": v} for k, v in
                   sorted(hour_buckets.items())]

    # Build time-series per alarm type (top 10)
    top_10_names = [name for name, _ in alarm_name_dist.most_common(10)]
    alarm_ts: dict[str, dict[str, int]] = {name: {} for name in top_10_names}
    for a in alarms:
        name = a.get("告警名称")
        t = a.get("发生时间")
        if name in alarm_ts and t:
            key = t.strftime("%Y-%m-%dT%H:00:00")
            alarm_ts[name][key] = alarm_ts[name].get(key, 0) + 1
    alarm_time_series = {
        name: [{"time": k, "count": v} for k, v in sorted(buckets.items())]
        for name, buckets in alarm_ts.items()
    }

    # Include results-level network managers
    nm_with_results = store.get_network_managers()
    meta_nms = meta.get("network_managers", [])
    all_nms = sorted(set(nm_with_results) | set(meta_nms))

    return StatsResponse(
        network_managers=all_nms,
        total_alarms=meta["total_records"],
        total_nes=meta["unique_ne"],
        total_alarm_types=meta["unique_alarm_names"],
        severity_distribution=severity_dist,
        top_alarm_names=[{"name": k, "count": v} for k, v in alarm_name_dist.most_common(20)],
        top_network_elements=[{"name": k, "count": v} for k, v in ne_dist.most_common(20)],
        railway_distribution=railway_dist,
        time_range={"min": time_min.isoformat() if time_min else None,
                     "max": time_max.isoformat() if time_max else None},
        time_series=time_series,
        alarm_time_series=alarm_time_series,
    )


@router.get("/export-alarms")
async def export_alarms(
    limit: int = 10000,
    network_manager: str = "",
):
    """Export alarm data as Excel file."""
    alarms = store.get_alarms()
    if not alarms:
        raise HTTPException(status_code=400, detail="无告警数据")

    if network_manager:
        alarms = [a for a in alarms if a.get("网管", "") == network_manager]
    if len(alarms) > limit:
        alarms = alarms[:limit]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "历史告警数据"

    headers = ["专业", "网管", "网元", "告警对象", "告警级别", "告警名称", "告警类型",
               "告警描述", "发生时间", "恢复时间", "铁路线", "站点", "机房", "厂商"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)

    for row_idx, a in enumerate(alarms, 2):
        for col_idx, h in enumerate(headers, 1):
            val = a.get(h, "")
            if isinstance(val, datetime):
                val = val.isoformat()
            ws.cell(row=row_idx, column=col_idx, value=val)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    wb.close()

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=alarm_data.xlsx"},
    )


@router.get("/alarm-preview")
async def alarm_preview(
    page: int = 1,
    page_size: int = 100,
    network_manager: str = "",
    severity: str = "",
    search: str = "",
):
    """Get alarm data preview for Dashboard table."""
    alarms = store.get_alarms()
    if not alarms:
        return {"rows": [], "total": 0}

    if network_manager:
        alarms = [a for a in alarms if a.get("网管", "") == network_manager]
    if severity:
        alarms = [a for a in alarms if a.get("告警级别", "") == severity]
    if search:
        q = search.lower()
        alarms = [a for a in alarms if q in str(a.get("告警名称", "")).lower()
                  or q in str(a.get("网元", "")).lower()]

    total = len(alarms)
    start = (page - 1) * page_size
    rows = []
    for a in alarms[start:start + page_size]:
        rows.append({
            "ne_name": a.get("网元", ""), "alarm_name": a.get("告警名称", ""),
            "severity": a.get("告警级别", ""), "network_manager": a.get("网管", ""),
            "alarm_object": a.get("告警对象", ""), "alarm_type": a.get("告警类型", ""),
            "alarm_desc": a.get("告警描述", ""),
            "occur_time": a.get("发生时间").isoformat() if a.get("发生时间") else "",
            "resume_time": a.get("恢复时间").isoformat() if a.get("恢复时间") else "",
            "railway": a.get("铁路线", ""), "station": a.get("站点", ""),
        })

    return {"rows": rows, "total": total}
